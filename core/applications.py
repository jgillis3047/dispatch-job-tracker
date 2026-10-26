"""Applications you've sent: Excel import, email status matching, focus areas, geocoding, analytics."""
import hashlib
import json
import re
import threading
import time
from collections import Counter, defaultdict
from datetime import datetime, date, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse

from . import store

STATUS_LABEL = {"pending": "Pending", "interview": "Interviewing", "offer": "Offer", "rejected": "Not selected", "no_response": "No response"}
RANK = {"pending": 0, "no_response": 1, "rejected": 2, "interview": 3, "offer": 4}
EVENT_STATUS = {"assessment": "interview", "withdrawn": "rejected", "interview": "interview", "offer": "offer", "rejected": "rejected"}
_cfg_cache = {"t": 0, "v": None}


def _cfg():
    if time.time() - _cfg_cache["t"] > 5:
        _cfg_cache.update(t=time.time(), v=store.load_settings())
    return _cfg_cache["v"]

SCHEMA = """
CREATE TABLE IF NOT EXISTS applications (
  id TEXT PRIMARY KEY, data TEXT NOT NULL, updated TEXT
);
CREATE TABLE IF NOT EXISTS status_emails (message_id TEXT PRIMARY KEY, app_id TEXT, kind TEXT, seen TEXT);
"""


def db():
    con = store.db()
    con.executescript(SCHEMA)
    return con


# ---------------------------------------------------------------- normalizing
_SUFFIX = r"\b(?:inc|llc|l\.l\.c|ltd|corp|corporation|co|company|the|group|holdings|plc|lp|llp|na|n\.a|usa|us|of america|incorporated)\b"


def norm_company(s):
    s = (s or "").lower().replace("&", " and ")
    s = re.sub(r"[^\w\s]", " ", s)
    s = re.sub(_SUFFIX, " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_title(s):
    s = (s or "").lower()
    s = re.sub(r"\(.*?\)|\[.*?\]", " ", s)
    s = re.sub(r"\b(sr|senior|jr|junior|i{1,3}|iv|v|lead|principal|staff|level \d|\d)\b", " ", s)
    s = re.sub(r"[^\w\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def app_key(company, title):
    base = norm_company(company) + "|" + re.sub(r"\s+", " ", (title or "").lower().strip())
    return "app:" + hashlib.sha1(base.encode()).hexdigest()[:16]


FOCUS = [
    ("HRIS & HR tech", r"\bhris\b|hr systems|human resources? (?:information|systems|technology)|workday|peoplesoft|payroll systems|hcm"),
    ("Data engineering", r"data engineer|etl|data platform|data pipeline|big data|data architect|analytics engineer|data integration|integration engineer"),
    ("BI & reporting", r"\bbi\b|business intelligence|reporting|tableau|power bi|dashboard|visuali[sz]ation"),
    ("Data analysis", r"data analy|analytics|analyst|data scien|insights|quantitative|statistic|research scien"),
    ("Database & SQL", r"database|\bsql\b|dba|\becm\b|data (?:management|governance|quality|specialist|systems|coordinator|services)|master data|\bdata\b"),
    ("Software & apps", r"software|developer|full[- ]?stack|front[- ]?end|back[- ]?end|\bweb\b|application(?:s)? (?:engineer|developer)|programmer|mobile|\bai\b|machine learning|hardware engineer"),
    ("Product & project", r"product|project|program manager|scrum|business analyst|business systems"),
    ("IT & systems", r"\bit\b|information technology|technology|technician|help ?desk|service desk|support|systems?|network|infrastructure|desktop|cloud|devops|computer"),
    ("Management", r"manager|director|head of|team lead|supervisor|chief"),
]
_FOCUS_RX = [(n, re.compile(p, re.I)) for n, p in FOCUS]


def focus_area(title):
    t = title or ""
    for name, rx in _FOCUS_RX:
        if rx.search(t):
            return name
    return "Other"


PLATFORMS = [("Workday", r"myworkday"), ("Indeed", r"indeed"), ("LinkedIn", r"linkedin"), ("iCIMS", r"icims"),
             ("Greenhouse", r"greenhouse"), ("Lever", r"lever\.co"), ("SmartRecruiters", r"smartrecruiters"), ("Ashby", r"ashby"),
             ("Taleo / Oracle", r"taleo|oraclecloud"), ("SuccessFactors", r"successfactors|sapsf"), ("Paycom", r"paycom"),
             ("PeopleAdmin", r"peopleadmin"), ("UKG", r"ultipro|ukg"), ("ADP", r"adp\.com"), ("Gov jobs", r"governmentjobs|usajobs|neogov"),
             ("Jobvite", r"jobvite"), ("ZipRecruiter", r"ziprecruiter"), ("Glassdoor", r"glassdoor")]


def platform(link):
    l = (link or "").lower()
    if not l.startswith("http"):
        return "Indeed" if "indeed" in l else ("LinkedIn" if "linkedin" in l else "Not recorded")
    for name, p in PLATFORMS:
        if re.search(p, l):
            return name
    return "Company site"


# ---------------------------------------------------------------- storage
def _row(r):
    d = json.loads(r["data"])
    d["id"] = r["id"]
    return d


def all_apps():
    with db() as con:
        return [_row(r) for r in con.execute("SELECT * FROM applications").fetchall()]


def get_app(aid):
    with db() as con:
        r = con.execute("SELECT * FROM applications WHERE id=?", (aid,)).fetchone()
    return _row(r) if r else None


def save_app(a):
    a = dict(a)
    aid = a.pop("id")
    with store._lock, db() as con:
        con.execute("INSERT OR REPLACE INTO applications VALUES (?,?,?)", (aid, json.dumps(a, default=str), store.now()))


def recompute(a):
    """Derive status from the starting status plus employer emails, then apply the automatic rules."""
    cfg = _cfg()
    base = a.get("base_status") or "pending"
    status, status_date = base, a.get("applied_on")
    if status == "withdrawn":
        status = "rejected"
    for e in sorted(a.get("events", []), key=lambda e: e.get("date") or ""):
        k = EVENT_STATUS.get(e["kind"])
        if k:  # the latest employer email wins
            status, status_date = k, (e.get("date") or "")[:10] or status_date
    replies = [e for e in a.get("events", []) if e["kind"] in EVENT_STATUS and e.get("date")]
    age = None
    if a.get("applied_on"):
        try:
            age = (date.today() - date.fromisoformat(a["applied_on"][:10])).days
        except ValueError:
            pass
    a["age_days"] = age
    ghost = int(cfg.get("ghost_days") or 0)
    closed = None
    if status == "pending" and not replies and ghost and age is not None and age >= ghost:
        status = "no_response"   # automatic: nobody replied in `ghost` days
        closed = (date.fromisoformat(a["applied_on"][:10]) + timedelta(days=ghost)).isoformat()
        status_date = closed
    if base == "offer" or any(e["kind"] == "offer" for e in a.get("events", [])):
        status = "offer"  # an offer stays an offer unless you change it yourself
        closed = None
    if a.get("manual_status"):
        status = a["manual_status"]
    a["status"] = status
    a["status_label"] = STATUS_LABEL.get(status, status)
    a["status_date"] = status_date
    if status in ("rejected", "offer"):
        closed = next((e["date"][:10] for e in sorted(a.get("events", []), key=lambda e: e.get("date") or "", reverse=True)
                       if EVENT_STATUS.get(e["kind"]) in ("rejected", "offer") and e.get("date")), None)
    a["closed_on"] = closed
    a["confirmed"] = a.get("confirmed") or any(e["kind"] == "received" for e in a.get("events", []))
    if replies and a.get("applied_on"):
        first = min(e["date"][:10] for e in replies)
        try:
            a["days_to_reply"] = max(0, (date.fromisoformat(first) - date.fromisoformat(a["applied_on"][:10])).days)
        except ValueError:
            a["days_to_reply"] = None
    else:
        a["days_to_reply"] = None
    fu = int(cfg.get("followup_days") or 0)
    last_touch = max([x for x in (a.get("followed_up"), a.get("applied_on")) if x] or [""])
    try:
        since = (date.today() - date.fromisoformat(last_touch[:10])).days if last_touch else None
    except ValueError:
        since = None
    a["follow_up"] = bool(fu and status == "pending" and since is not None and since >= fu)
    a["focus"] = focus_area(a.get("title"))
    from .target import employer_type
    a["employer_type"] = employer_type(a.get("company"), a.get("title"))
    c, st = clean_place(a.get("city"), a.get("state"))
    a["place"] = ", ".join(x for x in (c, st) if x) or a.get("location") or ""
    a["platform"] = platform(a.get("link"))
    return a


def upsert_app(fields, source):
    aid = fields.get("id") or app_key(fields.get("company"), fields.get("title"))
    a = get_app(aid) or {"id": aid, "events": [], "source": source, "created": store.now()}
    for k, v in fields.items():
        if v not in (None, "") and (k not in a or source == "history" or not a.get(k)):
            a[k] = v
    a["id"] = aid
    save_app(recompute(a))
    return aid


# ---------------------------------------------------------------- excel
HEADERS = {
    "title": r"^(job|job title|title|position|role)$", "company": r"^(company|employer|organization|org)$",
    "city": r"^city$", "state": r"^(state|st)$", "location": r"^location$", "posted": r"^(date posted|posted|posted on|date opened|opened)$",
    "interview": r"^interview\??$", "offer": r"^offer\??$", "email": r"^email received$", "applied": r"^(applied on|date applied|applied|application date)$",
    "link": r"^(link|url|job link|posting)$", "notes": r"^(comments?|notes?)$", "status": r"^(status|result|outcome)$",
    "closed": r"^(date closed|closed|closed on|rejected on)$",
}


def find_excel(cfg):
    p = cfg.get("excel_path")
    if p and Path(p).exists():
        return Path(p)
    folder = store.APP_DIR.parent
    xs = sorted([x for x in folder.glob("*.xlsx") if not x.name.startswith("~$") and "dispatch" not in x.name.lower()],
                key=lambda x: ("job" not in x.name.lower() and "appl" not in x.name.lower(), -x.stat().st_mtime))
    return xs[0] if xs else None


def _d(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str) and v.strip():
        from .extract import norm_date
        return norm_date(v.strip())
    return None


def _yn(v):
    s = str(v or "").strip().lower()
    return "yes" if s in ("yes", "y", "true", "1") else "no" if s in ("no", "n", "false", "0") else "pending" if s.startswith("pend") else ""


def import_history(path=None):
    """One-time import of an old tracking spreadsheet. Dispatch never re-reads it afterwards."""
    path = Path(path) if path else find_excel(store.load_settings())
    if not path or not path.exists():
        return {"imported": 0, "file": None}
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    count = 0
    seen = set()
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = None
        for r in rows:
            cells = [str(c).strip().lower() if c is not None else "" for c in r]
            m = {}
            for key, pat in HEADERS.items():
                for i, c in enumerate(cells):
                    if re.match(pat, c):
                        m.setdefault(key, i)
            if "title" in m and ("company" in m or "applied" in m):
                header = m
                break
        if not header:
            continue
        for i, r in enumerate(rows, start=2):
            g = lambda k: r[header[k]] if k in header and header[k] < len(r) else None
            title = str(g("title") or "").strip()
            company = str(g("company") or "").strip()
            if not title and not company:
                continue
            interview, offer, status = _yn(g("interview")), _yn(g("offer")), str(g("status") or "").strip().lower()
            if offer == "yes" or "offer" in status:
                base = "offer"
            elif interview == "yes" or "interview" in status:
                base = "interview"
            elif re.search(r"reject|denied|declin|not selected|no$", status) or (interview == "no" and offer in ("no", "")):
                base = "rejected"
            else:
                base = "pending"
            city, state = str(g("city") or "").strip(), str(g("state") or "").strip()
            loc = str(g("location") or "").strip() or ", ".join(x for x in (city, state) if x)
            link = str(g("link") or "").strip()
            fields = {"title": title or "(role not recorded)", "company": company or "(company not recorded)", "city": city, "state": state,
                      "location": loc, "posted_on": _d(g("posted")), "applied_on": _d(g("applied")), "link": link,
                      "notes_excel": str(g("notes") or "").strip(), "base_status": base,
                      "confirmed": _yn(g("email")) == "yes", "excel_row": i, "excel_sheet": ws.title}
            if _d(g("closed")):
                fields["closed_excel"] = _d(g("closed"))
            aid = app_key(fields["company"], fields["title"])
            if aid in seen:
                aid = aid + f"-{i}"
            seen.add(aid)
            fields["id"] = aid
            fields["history"] = True
            upsert_app(fields, "history")
            count += 1
    store.set_state(history_file=path.name, history_imported=count, history_at=store.now())
    return {"imported": count, "file": path.name}


def reset_email_matches():
    """Forget email-derived data so employer emails are re-read with a better reader. History rows are kept."""
    with store._lock, db() as con:
        for r in con.execute("SELECT * FROM applications").fetchall():
            a = _row(r)
            if a.get("history") or a.get("source") == "dispatch":
                a["events"] = [e for e in a.get("events", []) if e.get("kind") == "followed_up"]
                d = dict(recompute(a)); d.pop("id")
                con.execute("UPDATE applications SET data=? WHERE id=?", (json.dumps(d, default=str), a["id"]))
            else:
                con.execute("DELETE FROM applications WHERE id=?", (a["id"],))
        con.execute("DELETE FROM status_emails")


def remove_history():
    """Delete imported spreadsheet rows. Rows that employer emails were matched to are kept as email-tracked applications."""
    removed = kept = 0
    with store._lock, db() as con:
        for r in con.execute("SELECT * FROM applications").fetchall():
            a = _row(r)
            if not a.get("history"):
                continue
            if a.get("events"):
                a["history"] = False
                a["source"] = "email"
                a.pop("base_status", None)
                a.pop("excel_row", None)
                d = dict(recompute(a)); d.pop("id")
                con.execute("UPDATE applications SET data=? WHERE id=?", (json.dumps(d, default=str), a["id"]))
                kept += 1
            else:
                con.execute("DELETE FROM applications WHERE id=?", (a["id"],))
                removed += 1
    store.set_state(history_imported=0, history_removed_at=store.now())
    return {"removed": removed, "kept": kept}


def history_offer():
    st = store.get_state()
    if st.get("history_imported"):
        return {"imported": st["history_imported"], "file": st.get("history_file")}
    x = find_excel(store.load_settings())
    return {"available": x.name if x else None}


# ---------------------------------------------------------------- email matching
def status_email_seen(mid):
    with db() as con:
        return con.execute("SELECT 1 FROM status_emails WHERE message_id=?", (mid,)).fetchone() is not None


def _word_in(needle, hay):
    if not needle or len(needle) < 2:
        return False
    return re.search(r"(?<![a-z0-9])" + re.escape(needle) + r"(?![a-z0-9])", hay) is not None


def match_event(ev, apps):
    hay = norm_company(f"{ev.get('sender_name')} {ev.get('sender')} {ev.get('subject')} {ev.get('text', '')[:4000]}")
    hay_dom = (ev.get("sender") or "").split("@")[-1].lower().replace("-", "")
    subj_title = norm_title(f"{ev.get('subject')} {ev.get('title_guess') or ''} {ev.get('text', '')[:1500]}")
    cands = []
    guess = norm_company(ev.get("company_guess"))
    for a in apps:
        c = norm_company(a.get("company"))
        if not c or c.startswith("company not recorded"):
            continue
        squashed = c.replace(" ", "")
        hit = _word_in(c, hay) or (len(squashed) >= 4 and squashed in hay_dom) or (guess and (guess == c or (len(guess) > 4 and (guess in c or c in guess))))
        if not hit:
            continue
        ev_d = (ev.get("date") or "")[:10]
        if ev_d and a.get("applied_on"):
            try:
                if date.fromisoformat(a["applied_on"][:10]) > date.fromisoformat(ev_d) + timedelta(days=2):
                    continue  # this email is older than the application, so it's about an earlier one
            except ValueError:
                pass
        tt = set(norm_title(a.get("title")).split())
        score = len(tt & set(subj_title.split())) / max(1, len(tt))
        cands.append((score, a.get("applied_on") or "", a))
    if not cands and ev.get("title_guess"):  # no company in the email: match on the exact job title instead
        tg = norm_title(ev["title_guess"])
        same = [a for a in apps if norm_title(a.get("title")) == tg and (a.get("applied_on") or "") <= ((ev.get("date") or "")[:10] or "9999")]
        if len(same) == 1:
            return same[0]
    if not cands:
        return None
    ev_date = (ev.get("date") or "")[:10]
    cands.sort(key=lambda x: (x[0], x[1] <= ev_date if ev_date else True, x[1]), reverse=True)
    return cands[0][2]


JOB_STATUS_FOR = {"received": "applied", "interview": "interviewing", "assessment": "interviewing", "offer": "offer",
                  "rejected": "passed", "withdrawn": "passed"}


def match_job(company, title):
    """Find the Dispatch job (from an alert or Radar) an application is about."""
    c = norm_company(company)
    if not c or c.startswith("company not") :
        return None
    tt = set(norm_title(title).split()) if title and not title.startswith("(") else set()
    best = None
    for j in store.all_jobs():
        jc = norm_company(j.get("company"))
        if not jc or not (jc == c or (len(c) > 4 and (c in jc or jc in c))):
            continue
        jt = set(norm_title(j.get("title")).split())
        score = len(tt & jt) / max(1, len(tt | jt)) if tt else 0.01
        if (tt and score >= 0.5 or not tt) and (best is None or score > best[0]):
            best = (score, j)
    return best[1] if best else None


def link_job(a, kind=None):
    """Copy posting details onto the application and move the job along the pipeline."""
    j = store.get_job(a["job_id"]) if a.get("job_id") else match_job(a.get("company"), a.get("title"))
    if not j:
        return a
    a["job_id"] = j["id"]
    jt = j.get("title") or ""
    if not a.get("title") or a["title"].startswith("(") or (jt.lower().startswith(a["title"].lower()) and len(jt) > len(a["title"])):
        a["title"] = jt
    for src, dst in (("posted_date", "posted_on"), ("location", "location"), ("apply_url", "link"), ("salary_text", "salary_label")):
        if j.get(src) and not a.get(dst):
            a[dst] = j[src]
    if not a.get("city") and j.get("location"):
        m = re.match(r"^\s*([^,(]+?),\s*([A-Z]{2})\b", j["location"])
        if m:
            a["city"], a["state"] = m.group(1), m.group(2)
        elif "remote" in j["location"].lower():
            a["city"] = "Remote"
    target = JOB_STATUS_FOR.get(kind or "received")
    order = ["discovered", "new", "saved", "applied", "interviewing", "offer", "passed"]
    if target and j.get("status") != "hidden" and (j.get("status") not in order or order.index(target) > order.index(j.get("status"))):
        store.update_meta(j["id"], status=target)
    return a


def apply_event(ev):
    """Attach a parsed status email to an application (creating one if needed). Returns (app_id, changed_status)."""
    apps = all_apps()
    a = match_event(ev, apps)
    created = False
    if a is None:
        if not ev.get("company_guess"):
            with store._lock, db() as con:
                con.execute("INSERT OR REPLACE INTO status_emails VALUES (?,?,?,?)", (ev["message_id"], None, ev["kind"], store.now()))
            return None, None
        fields = {"company": ev["company_guess"], "title": ev.get("title_guess") or "(role not in email)",
                  "applied_on": (ev.get("date") or "")[:10] or None, "base_status": "pending", "from_email": True}
        aid = upsert_app(fields, "email")
        a = get_app(aid)
        created = True
    before = a.get("status")
    link_job(a, ev["kind"])
    evs = a.setdefault("events", [])
    if not any(e.get("message_id") == ev["message_id"] for e in evs):
        evs.append({k: ev.get(k) for k in ("message_id", "kind", "date", "subject", "sender", "snippet")})
    recompute(a)
    save_app(a)
    with store._lock, db() as con:
        con.execute("INSERT OR REPLACE INTO status_emails VALUES (?,?,?,?)", (ev["message_id"], a["id"], ev["kind"], store.now()))
    changed = a["status"] if (a["status"] != before or created) and a["status"] != "pending" else None
    return a["id"], changed


def set_manual(aid, status=None, notes=None, followed_up=None):
    a = get_app(aid)
    if not a:
        return None
    if followed_up is not None:
        a["followed_up"] = date.today().isoformat() if followed_up else None
        a.setdefault("events", [])
        if followed_up:
            a["events"].append({"message_id": f"followup-{a['id']}-{date.today()}", "kind": "followed_up", "date": date.today().isoformat(),
                                "subject": "You followed up", "sender": "", "snippet": ""})
    if status is not None:
        a["manual_status"] = status or None
    if notes is not None:
        a["notes"] = notes
    save_app(recompute(a))
    return a


def from_job(job):
    """User marked a Dispatch job as applied -> make it an application."""
    loc = job.get("location") or ""
    m = re.match(r"^\s*([^,(]+?),\s*([A-Z]{2})\b", loc)
    fields = {"title": job.get("title"), "company": job.get("company") or "(company not recorded)", "location": loc,
              "city": m.group(1) if m else ("Remote" if "remote" in loc.lower() else ""), "state": m.group(2) if m else "",
              "posted_on": job.get("posted_date"), "applied_on": date.today().isoformat(), "link": job.get("apply_url") or job.get("url"),
              "base_status": "pending", "job_id": job.get("id"), "salary_label": job.get("salary_label")}
    return upsert_app(fields, "dispatch")


def applied_to(job):
    """Is this alert/Radar job one you've already applied to?"""
    c, t = norm_company(job.get("company")), set(norm_title(job.get("title")).split())
    if not c or not t:
        return None
    for a in all_apps():
        if norm_company(a.get("company")) == c:
            at = set(norm_title(a.get("title")).split())
            if at and (len(t & at) / len(t | at) >= 0.6 or at <= t or t <= at):
                return a
    return None


def weekly_summary():
    apps = all_apps()
    start = (date.today() - timedelta(days=7)).isoformat()
    sent = [a for a in apps if (a.get("applied_on") or "") >= start]
    evs = [(e, a) for a in apps for e in a.get("events", []) if (e.get("date") or "")[:10] >= start and e["kind"] in EVENT_STATUS]
    kinds = Counter(EVENT_STATUS[e["kind"]] for e, _ in evs)
    total = len(apps)
    heard = sum(1 for a in apps if a["status"] not in ("pending", "no_response"))
    return {"sent": len(sent), "replies": len(evs), "interviews": kinds.get("interview", 0), "offers": kinds.get("offer", 0),
            "rejections": kinds.get("rejected", 0), "total": total, "response_rate": round(100 * heard / total) if total else 0,
            "updates": [{"status": EVENT_STATUS[e["kind"]], "title": a.get("title"), "company": a.get("company")} for e, a in evs][:15],
            "follow_ups": sorted([a for a in apps if a.get("follow_up")], key=lambda a: a.get("applied_on") or "")[:10],
            "by_focus": Counter(a["focus"] for a in sent).most_common(3)}


# ---------------------------------------------------------------- geocoding
GEO_PATH = store.DATA_DIR / "geocode.json"
_geo_lock = threading.Lock()
_geo_running = {"on": False}
KNOWN = {  # instant coordinates for common cities; the rest are looked up once on OpenStreetMap
    "las vegas,nv": (36.17, -115.14), "henderson,nv": (36.04, -114.98), "north las vegas,nv": (36.2, -115.12), "reno,nv": (39.53, -119.81),
    "seattle,wa": (47.61, -122.33), "bellevue,wa": (47.61, -122.2), "redmond,wa": (47.67, -122.12), "renton,wa": (47.48, -122.21),
    "tacoma,wa": (47.25, -122.44), "kent,wa": (47.38, -122.23), "everett,wa": (47.98, -122.2), "kirkland,wa": (47.68, -122.21),
    "san francisco,ca": (37.77, -122.42), "san jose,ca": (37.34, -121.89), "oakland,ca": (37.8, -122.27), "los angeles,ca": (34.05, -118.24),
    "san diego,ca": (32.72, -117.16), "sacramento,ca": (38.58, -121.49), "irvine,ca": (33.68, -117.83), "fremont,ca": (37.55, -121.99),
    "dublin,ca": (37.7, -121.94), "santa clara,ca": (37.35, -121.96), "sunnyvale,ca": (37.37, -122.04), "palo alto,ca": (37.44, -122.14),
    "philadelphia,pa": (39.95, -75.17), "harrisburg,pa": (40.27, -76.88), "pittsburgh,pa": (40.44, -79.99), "hershey,pa": (40.29, -76.65),
    "new york,ny": (40.71, -74.01), "newark,nj": (40.74, -74.17), "jersey city,nj": (40.73, -74.08), "washington,dc": (38.91, -77.04),
    "arlington,va": (38.88, -77.1), "richmond,va": (37.54, -77.44), "baltimore,md": (39.29, -76.61), "wilmington,de": (39.74, -75.55),
    "chicago,il": (41.88, -87.63), "phoenix,az": (33.45, -112.07), "denver,co": (39.74, -104.99), "dallas,tx": (32.78, -96.8),
    "austin,tx": (30.27, -97.74), "houston,tx": (29.76, -95.37), "portland,or": (45.52, -122.68), "boston,ma": (42.36, -71.06),
    "atlanta,ga": (33.75, -84.39), "fort meade,md": (39.11, -76.73), "mclean,va": (38.93, -77.18), "vienna,va": (38.9, -77.27),
    "mountlake terrace,wa": (47.79, -122.31), "allentown,pa": (40.6, -75.49), "upper saddle river,nj": (41.06, -74.1),
    "washington,dc": (38.91, -77.04), "alexandria,va": (38.8, -77.05), "reston,va": (38.96, -77.36), "herndon,va": (38.97, -77.39),
    "chantilly,va": (38.89, -77.43), "bethesda,md": (38.98, -77.1), "rockville,md": (39.08, -77.15), "columbia,md": (39.2, -76.86),
    "annapolis junction,md": (39.12, -76.78), "silver spring,md": (39.0, -77.03), "suitland,md": (38.85, -76.92), "catonsville,md": (39.27, -76.73),
    "king of prussia,pa": (40.09, -75.38), "malvern,pa": (40.04, -75.51), "norristown,pa": (40.12, -75.34), "camden,nj": (39.93, -75.12),
    "hackensack,nj": (40.89, -74.04), "west orange,nj": (40.79, -74.24), "rutherford,nj": (40.83, -74.11), "paulsboro,nj": (39.83, -75.24),
    "burlington,nj": (40.07, -74.86), "newark,de": (39.68, -75.75), "dover,de": (39.16, -75.52), "ashburn,va": (39.04, -77.49),
    "lacey,wa": (47.03, -122.82), "olympia,wa": (47.04, -122.9), "lynnwood,wa": (47.82, -122.32), "bothell,wa": (47.76, -122.21),
    "summerlin,nv": (36.19, -115.3), "paradise,nv": (36.1, -115.15), "north wales,pa": (40.21, -75.28), "wilmington,pa": (41.0, -80.3), "detroit,mi": (42.33, -83.05), "salt lake city,ut": (40.76, -111.89),
}
STATE_CENTER = {"CA": (37.2, -119.5), "WA": (47.4, -120.5), "NV": (38.5, -117), "PA": (40.9, -77.8), "NJ": (40.1, -74.6), "VA": (37.5, -78.8),
                "MD": (39.0, -76.8), "DC": (38.91, -77.04), "DE": (39.0, -75.5), "NY": (42.9, -75.5), "IL": (40.0, -89.2), "MI": (44.3, -85.4),
                "TX": (31.0, -99.0), "AZ": (34.2, -111.7), "CO": (39.0, -105.5), "OR": (44.0, -120.5), "FL": (28.6, -82.4), "GA": (32.7, -83.4),
                "MA": (42.3, -71.8), "UT": (39.3, -111.7), "NC": (35.5, -79.4), "OH": (40.3, -82.8), "MN": (46.3, -94.3)}


def _geo_cache():
    try:
        return json.loads(GEO_PATH.read_text("utf-8"))
    except Exception:
        return {}


def _geo_fill(keys):
    if _geo_running["on"]:
        return
    _geo_running["on"] = True
    try:
        import requests
        cache = _geo_cache()
        for k in keys:
            if k in cache:
                continue
            city, st = k.split(",", 1)
            try:
                r = requests.get("https://nominatim.openstreetmap.org/search", timeout=12,
                                 params={"city": city, "state": st, "country": "USA", "format": "json", "limit": 1},
                                 headers={"User-Agent": "Dispatch-job-tracker/1.0 (personal use)"})
                res = r.json()
                cache[k] = [float(res[0]["lat"]), float(res[0]["lon"])] if res else None
            except Exception:
                break
            with _geo_lock:
                GEO_PATH.write_text(json.dumps(cache), "utf-8")
            time.sleep(1.1)  # OpenStreetMap's fair-use limit
    finally:
        _geo_running["on"] = False


def coords(city, state):
    k = f"{(city or '').strip().lower()},{(state or '').strip().lower()}"
    if k in KNOWN:
        return KNOWN[k], False
    c = _geo_cache().get(k)
    if c:
        return tuple(c), False
    return None, k not in _geo_cache()


# ---------------------------------------------------------------- analytics
US_STATES = {"alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA", "colorado": "CO", "connecticut": "CT",
    "delaware": "DE", "florida": "FL", "georgia": "GA", "hawaii": "HI", "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA",
    "kansas": "KS", "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD", "massachusetts": "MA", "michigan": "MI",
    "minnesota": "MN", "mississippi": "MS", "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV", "new hampshire": "NH",
    "new jersey": "NJ", "new mexico": "NM", "new york": "NY", "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK",
    "oregon": "OR", "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC", "south dakota": "SD", "tennessee": "TN",
    "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA", "west virginia": "WV", "wisconsin": "WI",
    "wyoming": "WY", "district of columbia": "DC", "dc": "DC", "d.c.": "DC", "dmv": "DC"}


def clean_place(city, state):
    """Tidy free-typed spreadsheet locations -> (City, ST)."""
    c, s = (city or "").strip(), (state or "").strip()
    sl = s.lower()
    st = US_STATES.get(sl, s.upper() if len(s) == 2 else s)
    cl = c.lower()
    if cl in ("remote", "anywhere") or sl == "remote":
        return "Remote", ""
    if cl in ("united states", "usa", "us") or (sl in ("usa", "us", "united states") and not c):
        return "Anywhere in US", ""
    if re.match(r"(?i)washington,? ?d\.?c\.?$", c) or cl in ("dc", "d.c."):
        return "Washington", "DC"
    if cl in US_STATES and (not st or US_STATES.get(cl) == st):  # state typed as the city
        return "", US_STATES[cl]
    return c.title() if c.isupper() or c.islower() else c, st


def _week(d):
    dd = date.fromisoformat(d[:10])
    return (dd - timedelta(days=dd.weekday())).isoformat()


def insights():
    apps = [a for a in all_apps() if not a.get("hidden")]
    total = len(apps)
    by_status = Counter(a["status"] for a in apps)
    heard = [a for a in apps if a["status"] not in ("pending", "no_response")]
    # cities
    cities = defaultdict(lambda: {"n": 0, "status": Counter()})
    need_geo = []
    for a in apps:
        city, st = clean_place(a.get("city"), a.get("state"))
        if not city and st:
            city = f"Somewhere in {st}"
        if not city and not st:
            city = "Not recorded"
        c = cities[(city, st)]
        c["n"] += 1
        c["status"][a["status"]] += 1
    city_rows = []
    for (city, st), v in cities.items():
        real = city not in ("Remote", "Not recorded", "Anywhere in US") and not city.startswith("Somewhere") and st
        pt, missing = coords(city, st) if real else (None, False)
        if missing:
            need_geo.append(f"{city.lower()},{st.lower()}")
        approx = False
        if not pt and st in STATE_CENTER and real:
            pt, approx = STATE_CENTER[st], True
        city_rows.append({"city": city, "state": st, "n": v["n"], "status": dict(v["status"]),
                          "lat": pt[0] if pt else None, "lng": pt[1] if pt else None, "approx": approx})
    city_rows.sort(key=lambda r: -r["n"])
    if need_geo:
        threading.Thread(target=_geo_fill, args=(need_geo,), daemon=True).start()
    states = Counter(clean_place(a.get("city"), a.get("state"))[1] or "—" for a in apps)
    focus = defaultdict(Counter)
    for a in apps:
        focus[a["focus"]][a["status"]] += 1
    # weekly
    weekly = defaultdict(lambda: {"applied": Counter(), "closed": 0, "replies": 0})
    for a in apps:
        if a.get("applied_on"):
            weekly[_week(a["applied_on"])]["applied"][a["status"]] += 1
        for e in a.get("events", []):
            if e.get("date") and e["kind"] in ("rejected", "withdrawn"):
                weekly[_week(e["date"])]["closed"] += 1
            if e.get("date") and e["kind"] in ("rejected", "withdrawn", "interview", "offer", "assessment"):
                weekly[_week(e["date"])]["replies"] += 1
        if a["status"] == "no_response" and a.get("closed_on"):
            weekly[_week(a["closed_on"])]["closed"] += 1
    if weekly:
        w0, w1 = date.fromisoformat(min(weekly)), date.fromisoformat(max(weekly))
        d = w0
        while d <= w1:
            weekly[d.isoformat()]
            d += timedelta(days=7)
    weeks = [{"week": k, "applied": dict(v["applied"]), "closed": v["closed"], "replies": v["replies"]} for k, v in sorted(weekly.items())]
    daily = Counter(a["applied_on"][:10] for a in apps if a.get("applied_on"))
    days = [a["days_to_reply"] for a in apps if a.get("days_to_reply") is not None]
    bins = [("0–3 days", 0, 3), ("4–7", 4, 7), ("8–14", 8, 14), ("15–30", 15, 30), ("31–60", 31, 60), ("60+", 61, 10**6)]
    hist = [{"label": l, "n": sum(1 for x in days if lo <= x <= hi)} for l, lo, hi in bins]
    plat = Counter(a["platform"] for a in apps)
    etype = defaultdict(Counter)
    from .target import employer_type
    for a in apps:
        etype[employer_type(a.get("company"), a.get("title"))][a["status"]] += 1
    lag = [(date.fromisoformat(a["applied_on"][:10]) - date.fromisoformat(a["posted_on"][:10])).days for a in apps
           if a.get("posted_on") and a.get("applied_on")]
    lag = [x for x in lag if 0 <= x < 365]
    applied_dates = sorted(daily)
    last7 = sum(n for d, n in daily.items() if d >= (date.today() - timedelta(days=6)).isoformat())
    pending_old = sum(1 for a in apps if a["status"] == "pending" and (a.get("age_days") or 0) >= 30)
    return {
        "total": total, "by_status": dict(by_status), "heard_back": len(heard),
        "response_rate": round(100 * len(heard) / total) if total else 0,
        "median_reply_days": sorted(days)[len(days) // 2] if days else None,
        "median_apply_lag": sorted(lag)[len(lag) // 2] if lag else None,
        "last7": last7, "pending_over_30": pending_old,
        "first_applied": applied_dates[0] if applied_dates else None, "last_applied": applied_dates[-1] if applied_dates else None,
        "cities": city_rows, "states": dict(states), "focus": {k: dict(v) for k, v in focus.items()}, "weeks": weeks,
        "daily": dict(daily), "reply_hist": hist, "platforms": dict(plat), "employer_types": {k: dict(v) for k, v in etype.items()}, "geo_pending": bool(need_geo),
        "with_email_updates": sum(1 for a in apps if a.get("events")),
        "follow_up": sum(1 for a in apps if a.get("follow_up")), "history": sum(1 for a in apps if a.get("history")),
        "ghost_days": int(_cfg().get("ghost_days") or 0),
    }
