"""Dispatch — job alert organizer, application tracker and job radar.
Run:  python app.py            (opens the app window)
      python app.py --background  (no window; used by "Start with Windows")"""
import csv
import io
import re
import os
import platform
import socket
import subprocess
import sys
import threading
import time
import webbrowser
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

HERE = Path(__file__).resolve().parent
os.chdir(HERE)
sys.path.insert(0, str(HERE))
if sys.stdout is None or sys.stderr is None:  # pythonw has no console
    (HERE / "data").mkdir(exist_ok=True)
    _log = open(HERE / "data" / "dispatch.log", "a", buffering=1, encoding="utf-8")
    sys.stdout = sys.stderr = _log

from flask import Flask, jsonify, request, send_from_directory, Response

from core import store, applications as A, radar as R, notify, target as T
from core.discover import search, search_links, core_query
from core.enrich import enrich
from core.extract import salary_label
from core.mail_reader import fetch, test_login, MailError
from core.parsers import parse_message, source_label
from core.status import parse_status_email
from core.demo import demo_jobs

PORT = 5057
URL = f"http://127.0.0.1:{PORT}"
app = Flask(__name__, static_folder="static", static_url_path="/static")

SYNC = {"running": False, "log": [], "started": None, "finished": None, "new": 0, "updates": 0, "error": None, "last_ok": None}
RADAR = {"running": False, "log": [], "error": None, "new": 0}
_sync_lock, _radar_lock = threading.Lock(), threading.Lock()


def log(msg, target=SYNC):
    target["log"] = (target["log"] + [msg])[-40:]
    print(time.strftime("%H:%M:%S"), msg, flush=True)


def decorate(job, resume_skills):
    job["source_label"] = source_label(job.get("source")) if job.get("origin") != "web" else ("Radar · " if job.get("radar") else "") + (job.get("board") or "Web")
    job["salary_label"] = salary_label(job.get("salary_min"), job.get("salary_max"), job.get("salary_period") or "year", job.get("salary_text"))
    job["match"] = store.match_score(job.get("skills"), resume_skills)
    rs = set(resume_skills)
    job["skills_have"] = [s for s in job.get("skills", []) if s in rs]
    job["skills_missing"] = [s for s in job.get("skills", []) if s not in rs]
    job["employer_type"] = job.get("employer_type") or T.employer_type(job.get("company"), job.get("title"), job.get("board"))
    return job


# ---------------------------------------------------------------- sync: alerts + application emails + excel
def run_sync(full=False):
    if not _sync_lock.acquire(blocking=False):
        return False
    SYNC.update(running=True, log=[], started=store.now(), finished=None, new=0, updates=0, error=None)
    changes, errors = [], []
    try:
        cfg = store.load_settings()
        accts = store.connected(cfg)
        if not accts:
            raise MailError("CONNECT: Connect your Gmail so Dispatch can read your job emails.")
        fresh = []
        for acct in accts:
            try:
                log(f"Signing in to {acct['email']}")
                # 1) job alerts
                for raw in fetch(acct, "alerts", cfg.get("days", 14), (cfg.get("extra_senders") or "").split(","), progress=log,
                                 seen=None if full else store.email_seen):
                    try:
                        meta, jobs = parse_message(raw)
                    except Exception as e:
                        continue
                    if store.email_seen(meta["message_id"]) and not full:
                        continue
                    for j in jobs:
                        j["seen_in"] = [meta["subject"]]
                        j["account"] = acct["email"]
                        if store.upsert_job(j, "email"):
                            fresh.append(j)
                    store.mark_email(meta, len(jobs))
                # 2) application updates (received / rejected / interview / offer)
                n = 0
                for raw in fetch(acct, "status", cfg.get("status_days", 150), progress=log, limit=1500,
                                 seen=None if full else A.status_email_seen):
                    try:
                        ev = parse_status_email(raw)
                    except Exception:
                        continue
                    if not ev or (A.status_email_seen(ev["message_id"]) and not full):
                        continue
                    aid, changed = A.apply_event(ev)
                    n += 1
                    if changed:
                        a = A.get_app(aid)
                        changes.append({"status": changed, "title": a.get("title"), "company": a.get("company"), "id": aid})
                if n:
                    log(f"{acct['email']}: {n} application updates read")
            except MailError as e:
                errors.append(str(e))
                log(str(e))
        SYNC["updates"] = len(changes)
        log(f"{len(fresh)} new postings from alerts. Opening each one for details…")
        SYNC["new"] = len(fresh)

        def work(j):
            enrich(j)
            store.upsert_job(j, "email")
            if cfg.get("auto_mark_applied", True) and A.applied_to(j):
                store.update_meta(j["id"], status="applied")  # you already applied: file it under Applied
            elif cfg.get("auto_filter_alerts", True):
                why = T.why_not(j, prof, lenient=True)
                if why:  # off-target alert (wrong kind of role, or outside your areas): keep it out of New
                    store.update_meta(j["id"], status="hidden")
                    store.upsert_job({"id": j["id"], "hidden_reason": why}, "email")
                    return False
            return True
        prof = T.profile(cfg)
        kept = 0
        with ThreadPoolExecutor(max_workers=6) as ex:
            for i, ok in enumerate(ex.map(work, fresh), 1):
                kept += bool(ok)
                if i % 3 == 0 or i == len(fresh):
                    log(f"Details pulled for {i} of {len(fresh)}")
        SYNC["new"] = kept
        if len(fresh) > kept:
            log(f"{len(fresh) - kept} alert postings didn't match your roles or areas and were filed under Hidden")
        if changes:
            prev = store.get_state().get("pending_changes", [])
            store.set_state(pending_changes=(prev + changes)[-60:])
            if store.get_state().get("initial_status_done"):
                notify.alert([], changes)
                store.set_state(pending_changes=[])
        store.set_state(initial_status_done=True)
        SYNC["last_ok"] = store.now()
        if errors and len(errors) == len(accts):
            SYNC["error"] = errors[0]
        log("Sync complete" + (f" ({len(errors)} account problem)" if errors else ""))
    except MailError as e:
        SYNC["error"] = str(e)
        log(str(e))
    except Exception as e:
        SYNC["error"] = f"Sync stopped: {type(e).__name__}: {e}"
        log(SYNC["error"])
    finally:
        SYNC.update(running=False, finished=store.now())
        _sync_lock.release()
    return True


def run_radar(send=True):
    if not _radar_lock.acquire(blocking=False):
        return
    RADAR.update(running=True, log=[], error=None, new=0)
    try:
        new, rep = R.run(progress=lambda m: log(m, RADAR))
        RADAR["new"] = len(new)
        if send and new:
            alert_jobs = sorted(new, key=lambda j: -(j.get("fit") or 0))
            res = notify.alert(alert_jobs, [])
            if res:
                log(f"Alert: email {res.get('email', 'off')}, desktop {res.get('desktop', 'off')}", RADAR)
    except Exception as e:
        RADAR["error"] = f"Radar stopped: {type(e).__name__}: {e}"
        log(RADAR["error"], RADAR)
    finally:
        RADAR["running"] = False
        _radar_lock.release()


def maybe_weekly_report(cfg, force=False):
    """Sunday evening progress email."""
    from datetime import datetime, date
    now = datetime.now()
    week = date.today().isocalendar()[:2]
    key = f"{week[0]}-W{week[1]}"
    if not force and (not cfg.get("weekly_report_on") or not store.connected(cfg) or now.weekday() != 6 or now.hour < 17
                      or store.get_state().get("weekly_sent") == key):
        return None
    summary = A.weekly_summary()
    rs = store.resume_profile(cfg)["skills"]
    finds = sorted([decorate(j, rs) for j in store.all_jobs() if j.get("radar") and j.get("status") == "discovered"],
                   key=lambda j: -(j.get("fit") or 0))[:5]
    res = notify.weekly(summary, finds)
    if res == "sent" and not force:
        store.set_state(weekly_sent=key)
    return res


CODE_FILES = [HERE / "app.py", *sorted((HERE / "core").glob("*.py"))]
_code_mtime = max(f.stat().st_mtime for f in CODE_FILES if f.exists())


def restart_if_updated():
    """When Dispatch's files are updated, restart into the new version on its own."""
    now = max(f.stat().st_mtime for f in CODE_FILES if f.exists())
    if now > _code_mtime + 1 and not SYNC["running"] and not RADAR["running"]:
        time.sleep(3)  # let a multi-file update finish copying
        log("Dispatch was updated. Restarting into the new version…")
        args = [sys.executable, str(HERE / "app.py"), "--restarted"] + [a for a in sys.argv[1:] if a in ("--background", "--no-browser")]
        if "--background" not in args:
            args.append("--no-browser")
        subprocess.Popen(args, cwd=str(HERE), creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        os._exit(0)


def first_run_setup():
    """Learn your targets from your past applications once, and clean out finds that don't fit."""
    st = store.get_state()
    if not st.get("history_imported") and not st.get("history_removed_at") and not A.all_apps():
        try:
            r = A.import_history()
            if r.get("imported"):
                log(f"Learned your target roles and areas from {r['imported']} past applications in {r['file']}")
        except Exception as e:
            log(f"Couldn't read your past applications ({type(e).__name__})")
    if store.get_state().get("alert_parser_version") != 2:
        try:
            with store.db() as con:
                bad = [r[0] for r in con.execute("SELECT id FROM jobs WHERE origin='email' AND status IN ('new','hidden') AND "
                                                  "(json_extract(data,'$.source') IN ('glassdoor','ziprecruiter','dice'))")]
                for jid in bad:
                    con.execute("DELETE FROM jobs WHERE id=?", (jid,))
                con.execute("DELETE FROM emails WHERE sender LIKE '%glassdoor%' OR sender LIKE '%ziprecruiter%' OR sender LIKE '%dice.com%'")
            store.set_state(alert_parser_version=2)
            log(f"Re-reading {len(bad)} Glassdoor, ZipRecruiter and Dice alert jobs with the improved reader…")
        except Exception as e:
            log(f"Couldn't reset alerts ({type(e).__name__})")
    if store.get_state().get("parser_version") != 5:
        try:
            A.reset_email_matches()
            store.set_state(parser_version=5, initial_status_done=False, pending_changes=[])
            log("Re-reading employer emails with the improved reader…")
        except Exception as e:
            log(f"Couldn't reset email matches ({type(e).__name__})")
    try:
        n = R.tidy_finds()
        if n:
            log(f"Hid {n} earlier Radar finds that don't match your roles or areas")
    except Exception:
        pass


def scheduler():
    time.sleep(4)
    first_run_setup()
    last_sync = last_radar = 0
    st = store.get_state()
    if st.get("radar_version") != 2:  # new, smarter Radar: search right away
        store.set_state(radar_version=2, radar_last=None)
        st = store.get_state()
    if st.get("radar_last"):
        try:
            from datetime import datetime
            last_radar = datetime.fromisoformat(st["radar_last"]).timestamp()
        except Exception:
            pass
    while True:
        cfg = store.load_settings()
        now = time.time()
        mins = int(cfg.get("auto_sync_minutes") or 0)
        if mins > 0 and store.connected(cfg) and now - last_sync >= mins * 60:
            run_sync()
            last_sync = time.time()
        restart_if_updated()
        maybe_weekly_report(cfg)
        hrs = float(cfg.get("radar_every_hours") or 0)
        if cfg.get("radar_enabled") and hrs > 0 and now - last_radar >= hrs * 3600:
            run_radar()
            last_radar = time.time()
        time.sleep(60)


# ---------------------------------------------------------------- routes: jobs
@app.get("/")
def index():
    return send_from_directory("static", "index.html")


@app.get("/api/jobs")
def api_jobs():
    resume = store.resume_profile()
    jobs = [decorate(j, resume["skills"]) for j in store.all_jobs()]
    st = store.get_state()
    return jsonify({"jobs": jobs, "resume": {"name": resume["name"], "skills": resume["skills"]},
                    "sync": SYNC, "radar": {**RADAR, "last": st.get("radar_last"), "found": st.get("radar_found")},
                    "accounts": [a["email"] for a in store.connected()],
                    "needs_connect": [a["email"] for a in store.load_settings()["accounts"] if not a.get("app_password")],
                    **store.email_stats()})


@app.patch("/api/jobs/<jid>")
def api_update(jid):
    b = request.get_json(force=True)
    store.update_meta(jid, b.get("status"), b.get("notes"))
    if b.get("status") == "applied":
        j = store.get_job(jid)
        if j:
            A.from_job(decorate(j, []))
    return jsonify(ok=True)


@app.post("/api/jobs/<jid>/refresh")
def api_refresh(jid):
    j = store.get_job(jid)
    if not j:
        return jsonify(error="Not found"), 404
    enrich(j)
    store.upsert_job(j, j.get("origin", "email"))
    return jsonify(decorate(store.get_job(jid), store.resume_profile()["skills"]))


@app.post("/api/jobs/<jid>/similar")
def api_similar(jid):
    j = store.get_job(jid)
    if not j:
        return jsonify(error="Not found"), 404
    cfg = store.load_settings()
    q = (request.get_json(silent=True) or {}).get("query") or core_query(j["title"])
    loc = j.get("location") or cfg.get("default_location", "")
    loc = "" if "remote" in (loc or "").lower() else loc.split("(")[0].strip()
    found, report = search(q, loc, cfg, ref_title=j["title"])
    rs = store.resume_profile()["skills"]
    for f in found:
        store.upsert_job(f, "web", similar_to=jid)
    return jsonify(query=q, location=loc, report=report, links=search_links(q, loc),
                   jobs=[decorate(dict(f, origin="web", status="discovered", similar_to=jid), rs) for f in found])


@app.post("/api/discover")
def api_discover():
    b = request.get_json(force=True)
    cfg = store.load_settings()
    q, loc = b.get("query", "").strip(), b.get("location", "").strip()
    if not q:
        return jsonify(error="Type a role to search for"), 400
    found, report = search(q, loc, cfg)
    rs = store.resume_profile()["skills"]
    for f in found:
        store.upsert_job(f, "web")
    return jsonify(query=q, location=loc, report=report, links=search_links(q, loc),
                   jobs=[decorate(dict(f, origin="web", status="discovered"), rs) for f in found])


@app.post("/api/sync")
def api_sync():
    full = bool((request.get_json(silent=True) or {}).get("full"))
    if SYNC["running"]:
        return jsonify(started=False, sync=SYNC)
    threading.Thread(target=run_sync, args=(full,), daemon=True).start()
    time.sleep(0.2)
    return jsonify(started=True, sync=SYNC)


@app.get("/api/sync")
def api_sync_status():
    return jsonify(SYNC)


# ---------------------------------------------------------------- routes: applications + insights
@app.get("/api/applications")
def api_apps():
    st = store.get_state()
    apps = sorted(A.all_apps(), key=lambda a: a.get("applied_on") or "", reverse=True)
    return jsonify(apps=apps, history=A.history_offer())


@app.patch("/api/applications/<aid>")
def api_app_update(aid):
    b = request.get_json(force=True)
    a = A.set_manual(aid, b.get("status") if "status" in b else None, b.get("notes"), b.get("followed_up"))
    return jsonify(a or {})


@app.get("/api/history")
def api_history():
    return jsonify(A.history_offer())


@app.post("/api/history")
def api_history_import():
    try:
        return jsonify(A.import_history())
    except Exception as e:
        return jsonify(error=f"Couldn't read the spreadsheet ({type(e).__name__}). If it's open in Excel, close it and try again."), 400


@app.delete("/api/history")
def api_history_remove():
    return jsonify(A.remove_history())


@app.post("/api/report/test")
def api_report_test():
    return jsonify(result=maybe_weekly_report(store.load_settings(), force=True))


@app.get("/api/insights")
def api_insights():
    return jsonify(A.insights())


@app.get("/api/applications/export.xlsx")
def api_export_apps():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    cols = [("Job", "title"), ("Company", "company"), ("City", "city"), ("State", "state"), ("Focus area", "focus"),
            ("Date Posted", "posted_on"), ("Applied On", "applied_on"), ("Status", "status_label"), ("Status date", "status_date"),
            ("Closed On", "closed_on"), ("Days to reply", "days_to_reply"), ("Applied via", "platform"), ("Link", "link"),
            ("Last email", "last_email"), ("Notes", "notes")]
    ws.append([c for c, _ in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="16213A")
    for a in sorted(A.all_apps(), key=lambda a: a.get("applied_on") or "", reverse=True):
        evs = sorted(a.get("events", []), key=lambda e: e.get("date") or "")
        a["last_email"] = f"{evs[-1]['date'][:10]}: {evs[-1]['subject']}" if evs else ""
        a["notes"] = " ".join(x for x in (a.get("notes_excel"), a.get("notes")) if x)
        if a.get("follow_up"):
            a["status_label"] += " (follow up)"
        ws.append([a.get(k) for _, k in cols])
    for col, w in zip("ABCDEFGHIJKLMNO", (42, 30, 16, 7, 18, 12, 12, 16, 12, 12, 9, 16, 50, 50, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=Dispatch applications.xlsx"})


# ---------------------------------------------------------------- routes: radar
@app.get("/api/radar")
def api_radar():
    cfg = store.load_settings()
    st = store.get_state()
    rs = store.resume_profile()["skills"]
    finds = [decorate(j, rs) for j in store.all_jobs() if j.get("radar")]
    boards = R.watched_boards(cfg)
    from core import boards as BD
    prof = R.profile(cfg)
    src = cfg.get("radar_sources") or {}
    names = ["LinkedIn", "LinkedIn (public sector & credit unions)"] + list(BD.SOURCES) + ["Company career sites", "Remote job boards"]
    notes = {"LinkedIn": "Every role in every area, past 24 hours", "LinkedIn (public sector & credit unions)": "Credit unions, universities, schools and government in your areas",
             "Company career sites": "Workday, Greenhouse, Lever, Ashby and SmartRecruiters sites of companies you've applied to",
             "Remote job boards": "Remotive, RemoteOK, The Muse", **{k: v["note"] for k, v in BD.SOURCES.items()}}
    # HigherEdJobs is left out on purpose: it blocks automated readers
    sources = [{"name": n, "on": src.get(n, True), "note": notes.get(n, "")} for n in names]
    if "CA" in prof["excluded_states"]:
        for x in sources:
            if x["name"] == "EdJoin":
                x["note"] += " (paused while California is blocked)"
    indeed = BD.indeed_alert_links(prof["roles"], [a["label"] for a in prof["areas"]], prof["remote"])
    alert_sites = BD.alert_links(prof["roles"], [a["label"] for a in prof["areas"]], prof["remote"])
    return jsonify(status=RADAR, profile=prof, defaults=R.default_profile(), last=st.get("radar_last"), sources=sources, indeed=indeed, alert_sites=alert_sites,
                   favor_types=cfg.get("radar_favor_types"), only_types=cfg.get("radar_only_types"),
                   report=st.get("radar_report"), found=st.get("radar_found"), sites=[{"label": k, "company": v[2]} for k, v in boards.items()],
                   last_alert=st.get("last_alert"), last_alert_result=st.get("last_alert_result"), last_alert_headline=st.get("last_alert_headline"),
                   finds=finds)


@app.post("/api/radar/run")
def api_radar_run():
    if RADAR["running"]:
        return jsonify(started=False, status=RADAR)
    threading.Thread(target=run_radar, daemon=True).start()
    time.sleep(0.2)
    return jsonify(started=True, status=RADAR)


@app.post("/api/alerts/test")
def api_alert_test():
    rs = store.resume_profile()["skills"]
    finds = sorted([decorate(j, rs) for j in store.all_jobs() if j.get("radar") or j.get("demo")], key=lambda j: -(j.get("match") or 0))[:5]
    if not finds:
        finds = [{"title": "Example: Data Analyst", "company": "Your next employer", "location": "Las Vegas, NV", "match": 88, "apply_url": URL}]
    res = notify.alert(finds, [], force_email=True)
    return jsonify(res)


# ---------------------------------------------------------------- routes: settings
@app.get("/api/settings")
def api_settings():
    r = store.resume_profile()
    return jsonify(settings=store.public_settings(), resume={"name": r["name"], "path": r["path"], "skills": r["skills"]},
                   history=A.history_offer(), windows=platform.system() == "Windows")


@app.post("/api/settings")
def api_save_settings():
    b = request.get_json(force=True)
    before = store.load_settings().get("start_with_windows")
    store.save_settings(b)
    if "start_with_windows" in b and bool(b["start_with_windows"]) != bool(before):
        startup_shortcut(bool(b["start_with_windows"]))
    return api_settings()


@app.post("/api/connect")
def api_connect():
    """One-step Gmail connection: check the code with Google, save it, start the first sync."""
    b = request.get_json(force=True)
    email = (b.get("email") or "").strip()
    pw = re.sub(r"\s+", "", b.get("app_password") or "")
    if not email or len(pw) < 12:
        return jsonify(ok=False, message="Paste the 16-letter code Google showed you.")
    try:
        test_login(email, pw)
    except MailError as e:
        return jsonify(ok=False, message=str(e))
    except Exception as e:
        return jsonify(ok=False, message=f"Couldn't reach Gmail ({type(e).__name__}). Check your internet connection.")
    cfg = store.load_settings()
    accts = [a for a in cfg["accounts"] if a["email"].lower() != email.lower()]
    accts.insert(0, {"email": email, "app_password": pw, "imap_host": ""})
    store.save_settings({"accounts": [dict(a) for a in accts]})
    if not SYNC["running"]:
        threading.Thread(target=run_sync, daemon=True).start()
    return jsonify(ok=True, message=f"Connected to {email}. Reading your job emails now…")


@app.post("/api/settings/test")
def api_test():
    b = request.get_json(force=True)
    pw = b.get("app_password")
    if not pw or pw == store.MASK:
        pw = (store.account_for(b.get("email")) or {}).get("app_password")
    try:
        test_login(b.get("email"), pw, b.get("imap_host") or None)
        return jsonify(ok=True, message="Connected.")
    except MailError as e:
        return jsonify(ok=False, message=str(e))
    except Exception as e:
        return jsonify(ok=False, message=f"{type(e).__name__}: {e}")


@app.post("/api/demo")
def api_demo():
    for j in demo_jobs():
        store.upsert_job(j, "email")
    return jsonify(ok=True)


@app.delete("/api/demo")
def api_demo_clear():
    store.delete_where(demo=True)
    return jsonify(ok=True)


@app.get("/api/export.csv")
def api_export():
    rs = store.resume_profile()["skills"]
    buf = io.StringIO()
    cols = ["status", "title", "company", "location", "work_style", "salary_label", "posted_date", "source_label",
            "match", "skills", "requirements", "apply_url", "notes"]
    w = csv.writer(buf)
    w.writerow(cols)
    for j in store.all_jobs():
        if j.get("status") == "discovered":
            continue
        j = decorate(j, rs)
        j["skills"] = ", ".join(j.get("skills", []))
        j["requirements"] = ", ".join(r["label"] + (" (optional)" if r.get("optional") else "") for r in j.get("requirements", []))
        w.writerow([j.get(c, "") for c in cols])
    return Response(buf.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=dispatch-jobs.csv"})


@app.post("/api/quit")
def api_quit():
    threading.Timer(0.5, lambda: os._exit(0)).start()
    return jsonify(ok=True)


def ui_version():
    files = [HERE / "static" / n for n in ("index.html", "app.js", "views.js", "app.css")]
    return str(int(max(f.stat().st_mtime for f in files if f.exists())))


@app.get("/api/diag/sources")
def api_diag_sources():
    """Quick check of how each job site answers Dispatch's requests."""
    from core.enrich import SESSION
    tests = {"HigherEdJobs": "https://www.higheredjobs.com/search/advanced_action.cfm?Keyword=Data+Analyst&OnlyTitle=1&SortBy=1&NumJobs=100",
             "Dice": "https://www.dice.com/jobs?q=%22Data+Analyst%22&location=Seattle%2C+WA%2C+USA&radius=50&radiusUnit=mi&filters.postedDate=SEVEN",
             "GovernmentJobs": "https://www.governmentjobs.com/jobs?keyword=Data+Analyst&location=Seattle%2C+WA&distance=50"}
    out = {}
    for k, u in tests.items():
        try:
            r = SESSION.get(u, timeout=20)
            t = r.text
            out[k] = {"status": r.status_code, "bytes": len(t), "title": (re.search(r"<title>(.*?)</title>", t, re.S | re.I) or [None, ""])[1].strip()[:80],
                      "hej_links": t.count("details.cfm?JobCode="), "dice_cards": t.count('data-testid="job-card"'), "server": r.headers.get("server")}
        except Exception as e:
            out[k] = {"error": f"{type(e).__name__}: {e}"[:200]}
    import traceback
    from core import boards as BD
    try:
        out["dice_parse"] = len(BD.dice(["Data Analyst", "Data Engineer"], "Seattle, WA"))
    except Exception:
        out["dice_parse"] = traceback.format_exc()[-1200:]
    try:
        r = SESSION.get(tests["HigherEdJobs"], timeout=20)
        out["hej_body"] = r.text[:900]
        out["hej_cookies"] = list(r.cookies.keys())
        out["hej_history"] = [h.status_code for h in r.history]
    except Exception as e:
        out["hej_body"] = str(e)
    return jsonify(out)


@app.get("/api/diag/alert")
def api_diag_alert():
    """Structure of the newest alert email from one sender (links and nearby text), to tune the alert reader."""
    import imaplib
    from email import policy
    from email.parser import BytesParser
    from bs4 import BeautifulSoup
    from core.mail_reader import connect, _select_all_mail, _gm_quote
    sender = request.args.get("from", "glassdoor.com")
    acct = store.connected()[0]
    M = connect(acct["email"], acct["app_password"])
    try:
        _select_all_mail(M, True)
        typ, data = M.uid("SEARCH", "X-GM-RAW", _gm_quote(f"from:{sender} newer_than:30d"))
        uids = data[0].split()
        if not uids:
            return jsonify(error="no email")
        typ, d = M.uid("FETCH", uids[-1], "(BODY.PEEK[])")
        msg = BytesParser(policy=policy.default).parsebytes(d[0][1])
    finally:
        M.logout()
    html = next((p.get_content() for p in msg.walk() if p.get_content_type() == "text/html"), "")
    text_part = next((p.get_content() for p in msg.walk() if p.get_content_type() == "text/plain"), "")
    if request.args.get("text"):
        return jsonify(subject=str(msg.get("Subject")), has_html=bool(html), text=re.sub(r"https?://\S+", "<URL>", text_part)[:3500])
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["style", "script", "head"]):
        t.decompose()
    links = []
    for a in soup.find_all("a", href=True)[:60]:
        box = a
        for _ in range(4):
            if box.parent is not None:
                box = box.parent
        links.append({"href": re.sub(r"\?.*", "", a["href"])[:120], "text": a.get_text(" ", strip=True)[:80],
                      "img_alt": [i.get("alt") for i in a.find_all("img")][:2],
                      "box": [x for x in box.get_text("\n", strip=True).split("\n")][:8]})
    meta, jobs = parse_message(d[0][1])
    return jsonify(subject=str(msg.get("Subject")), sender=str(msg.get("From")), links=links,
                   parsed=[{k: j.get(k) for k in ("title", "company", "location", "url")} for j in jobs[:10]])


@app.get("/api/ping")
def ping():
    return jsonify(app="dispatch", ui=ui_version(), code=_code_mtime)


# ---------------------------------------------------------------- windows integration
def _ps(script):
    subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script], timeout=30,
                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0), capture_output=True)


def make_shortcut(lnk_path, args):
    pyw = HERE / ".venv" / "Scripts" / "pythonw.exe"
    if not pyw.exists():
        pyw = Path(sys.executable).with_name("pythonw.exe")
    icon = HERE / "static" / "dispatch.ico"
    _ps(f"""$s=(New-Object -ComObject WScript.Shell).CreateShortcut('{lnk_path}');$s.TargetPath='{pyw}';$s.Arguments='"{HERE / 'app.py'}" {args}';
$s.WorkingDirectory='{HERE}';$s.IconLocation='{icon}';$s.Description='Dispatch job tracker';$s.Save()""")


def startup_shortcut(on):
    if platform.system() != "Windows":
        return
    lnk = Path(os.environ.get("APPDATA", "")) / "Microsoft" / "Windows" / "Start Menu" / "Programs" / "Startup" / "Dispatch.lnk"
    if on:
        make_shortcut(str(lnk), "--background")
    elif lnk.exists():
        lnk.unlink()


def ensure_desktop_shortcut():
    if platform.system() != "Windows":
        return
    try:
        desk = subprocess.run(["powershell", "-NoProfile", "-Command", "[Environment]::GetFolderPath('Desktop')"], capture_output=True,
                              text=True, timeout=20, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0)).stdout.strip()
        lnk = Path(desk) / "Dispatch.lnk"
        if desk and not lnk.exists():
            make_shortcut(str(lnk), "")
            for menu in [Path(os.environ.get("APPDATA", "")) / "Microsoft" / "Windows" / "Start Menu" / "Programs" / "Dispatch.lnk"]:
                make_shortcut(str(menu), "")
    except Exception as e:
        print("shortcut:", e)


def open_window():
    """Open Dispatch in its own app window (Edge or Chrome app mode), falling back to the default browser."""
    if platform.system() == "Windows":
        pf = [os.environ.get(k, "") for k in ("ProgramFiles(x86)", "ProgramFiles", "LOCALAPPDATA")]
        for exe in [Path(p) / "Microsoft" / "Edge" / "Application" / "msedge.exe" for p in pf] + \
                   [Path(p) / "Google" / "Chrome" / "Application" / "chrome.exe" for p in pf]:
            if exe.exists():
                subprocess.Popen([str(exe), f"--app={URL}", "--window-size=1440,960"])
                return
    webbrowser.open(URL)


def already_running():
    try:
        with socket.create_connection(("127.0.0.1", PORT), timeout=0.5):
            return True
    except OSError:
        return False


if __name__ == "__main__":
    background = "--background" in sys.argv
    if "--restarted" in sys.argv:  # wait for the old copy to let go of the port
        for _ in range(40):
            if not already_running():
                break
            time.sleep(0.5)
    if already_running():
        # Is the copy that's running older than these files? Then replace it with this one.
        outdated = True
        try:
            import json as _json
            from urllib.request import urlopen, Request
            info = _json.loads(urlopen(URL + "/api/ping", timeout=3).read())
            outdated = float(info.get("code") or 0) + 1 < _code_mtime
            if outdated:
                urlopen(Request(URL + "/api/quit", data=b"{}", headers={"Content-Type": "application/json"}), timeout=3)
        except Exception:
            pass
        if outdated:
            for _ in range(40):
                if not already_running():
                    break
                time.sleep(0.5)
        if already_running():
            if not background:
                open_window()
            sys.exit(0)
    threading.Thread(target=scheduler, daemon=True).start()
    if "--no-browser" not in sys.argv and not background:
        threading.Timer(1.2, open_window).start()
        threading.Thread(target=ensure_desktop_shortcut, daemon=True).start()
    print(f"\n  Dispatch is running at {URL}\n  You can close this window; use the Dispatch icon on your desktop to open it again.\n")
    app.run(host="127.0.0.1", port=PORT, debug=False, threaded=True)
